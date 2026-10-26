"""Local web console for running the CTA backtest.

Start from this directory with::

    python web_app.py

Then open http://127.0.0.1:8765.
"""

from __future__ import annotations

import json
import math
import os
import subprocess
import sys
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

from strategies import discover_strategy_catalog


BASE = Path(__file__).resolve().parent
OUTPUT = BASE / "output"
HOST = os.environ.get("CTA_WEB_HOST", "0.0.0.0")
PORT = 8765


HTML = r"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>CTA 回测控制台</title>
  <style>
    :root { color-scheme: light; font-family: -apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif; }
    body { margin: 0; background: #f3f5f7; color: #1f2937; }
    main { max-width: 1100px; margin: 32px auto; padding: 0 20px; }
    h1 { margin-bottom: 6px; } .hint { color: #6b7280; margin-top: 0; }
    form, section { background: white; border-radius: 12px; padding: 20px; margin: 16px 0; box-shadow: 0 2px 12px #0000000d; }
    .grid { display: grid; grid-template-columns: repeat(auto-fit,minmax(260px,1fr)); gap: 14px 18px; }
    label { display: flex; flex-direction: column; gap: 6px; font-weight: 600; font-size: 14px; }
    input, select { border: 1px solid #d1d5db; border-radius: 7px; padding: 9px 10px; font: inherit; font-weight: 400; }
    small { color: #6b7280; font-weight: 400; }
    .checks { display: flex; gap: 20px; margin-top: 18px; font-weight: 400; }
    .checks label { flex-direction: row; align-items: center; font-weight: 400; }
    button { margin-top: 20px; border: 0; border-radius: 8px; padding: 11px 18px; background: #2563eb; color: white; font-size: 15px; cursor: pointer; }
    button:disabled { background: #9ca3af; cursor: wait; }
    pre { white-space: pre-wrap; word-break: break-word; background: #111827; color: #e5e7eb; border-radius: 8px; padding: 16px; min-height: 180px; overflow: auto; }
    .status { font-weight: 600; } .ok { color: #047857; } .bad { color: #b91c1c; }
    .chart-tools { display:flex; gap:6px; margin: -8px 0 8px; }
    .chart-tools button { margin:0; padding:5px 10px; font-size:12px; background:#475569; }
    .chart-tip { min-height:20px; margin-top:6px; color:#334155; font:12px ui-monospace,Consolas,monospace; }
    .viewer-layout { display:grid; grid-template-columns:minmax(0,1fr) 180px; gap:16px; align-items:start; }
    .viewer-main { min-width:0; }
    .chart-grid { display:grid; grid-template-columns:repeat(2,minmax(0,1fr)); gap:16px; }
    .chart-grid section { min-width:0; }
    .viewer-side { background:#fff; border:1px solid #e5e7eb; border-radius:10px; padding:12px; position:sticky; top:12px; }
    .symbol-list { max-height:520px; overflow:auto; border:1px solid #d1d5db; border-radius:8px; outline:none; background:#fff; }
    .symbol-list:focus { box-shadow:0 0 0 2px #93c5fd; }
    .symbol-item { padding:8px 10px; cursor:pointer; border-bottom:1px solid #eef2f7; font-size:14px; }
    .symbol-item:last-child { border-bottom:0; }
    .symbol-item:hover { background:#eff6ff; }
    .symbol-item.active { background:#2563eb; color:#fff; }
    .strategy-picker { position:relative; }
    .strategy-toggle { width:100%; margin:0; text-align:left; background:#fff; color:#1f2937; border:1px solid #d1d5db; padding:9px 10px; font-size:14px; }
    .strategy-menu { position:absolute; z-index:20; left:0; right:0; top:calc(100% + 4px); max-height:280px; overflow:auto; background:#fff; border:1px solid #d1d5db; border-radius:8px; box-shadow:0 8px 24px #0002; padding:8px; }
    .strategy-menu[hidden] { display:none; }
    .strategy-option { display:flex; flex-direction:row; align-items:center; gap:8px; padding:7px 5px; font-weight:400; cursor:pointer; }
    .strategy-option:hover { background:#eff6ff; }
    .strategy-actions { display:flex; gap:8px; border-bottom:1px solid #e5e7eb; padding-bottom:7px; margin-bottom:4px; }
    .strategy-actions button { margin:0; padding:5px 8px; font-size:12px; background:#475569; }
    .strategy-categories { display:flex; flex-wrap:wrap; gap:6px; padding:4px 0 8px; border-bottom:1px solid #e5e7eb; margin-bottom:4px; }
    .strategy-category { margin:0; padding:5px 8px; font-size:12px; background:#e2e8f0; color:#334155; }
    .strategy-category.active { background:#2563eb; color:#fff; }
    @media (max-width:900px) { .viewer-layout { grid-template-columns:1fr; } .viewer-side { position:static; } }
  </style>
</head>
<body>
<main>
  <h1>CTA 回测控制台</h1>
  <p class="hint">页面参数对应 run_backtest.py 的命令行参数。数据路径填写包含 RB.h5、CU.h5 等文件的目录。</p>
  <form id="form">
    <div class="grid">
      <label>策略
        <div class="strategy-picker" id="strategyPicker">
          <button type="button" class="strategy-toggle" id="strategyToggle">加载策略...</button>
          <div class="strategy-menu" id="strategyMenu" hidden>
            <div class="strategy-actions"><button type="button" id="strategySelectAll">全选</button><button type="button" id="strategyClearAll">取消全选</button></div>
            <div class="strategy-categories" id="strategyCategories"></div>
            <label class="strategy-option"><input type="checkbox" id="strategyAll"> 全部策略</label>
            <div id="strategyOptions"></div>
          </div>
        </div>
        <input id="strategy" type="hidden" value="ma5ma10">
        <small>可选择一个、多个或全部策略</small>
      </label>
      <label>K线周期
        <select id="timeframe"><option value="1d">日线（1d）</option><option value="1h">1小时（1h）</option><option value="15m">15分钟（15m）</option></select>
        <small>传递给所有选中的策略</small>
      </label>
      <label>交易方向
        <select id="direction"><option value="both">做多和做空</option><option value="long">仅做多</option><option value="short">仅做空</option></select>
        <small>传递给所有选中的策略</small>
      </label>
      <label>品种
        <input id="symbols" value="RB">
        <small>多个品种用空格分隔；留空则按 universe.yaml 自动选择</small>
      </label>
      <label>组合层
        <select id="adapter"><option>passthrough</option><option>capped_capital</option><option>conflict_dedup</option></select>
      </label>
      <label>基础开仓手数
        <input id="volume" type="number" min="1" step="1" value="1">
        <small>填 1 使用固定基础 1 手；网格按每层计算，仍可累积多层。留空保留原配置。</small>
      </label>
      <label>品种级别
        <select id="level"><option value="product">product</option><option value="contract">contract</option></select>
      </label>
      <label>开始日期
        <input id="start" type="date" value="2024-01-01">
      </label>
      <label>结束日期
        <input id="end" type="date" value="2024-06-30">
      </label>
      <label>数据源
        <select id="dataSource"><option value="hdf5_index">hdf5_index</option><option value="live_archive">live_archive</option></select>
      </label>
      <label>HDF5 数据目录
        <input id="dataDir" value="../data/future_price" required>
        <small>相对于 strategy 目录，或填写绝对路径</small>
      </label>
      <label>输出前缀
        <input id="prefix" value="web_ma5ma10_rb" required>
      </label>
      <label>输出目录
        <input id="outDir" value="output">
      </label>
      <label>预热 bars
        <input id="warmupBars" type="number" min="0" value="500">
      </label>
      <label>配置目录
        <input id="configDir" value="config">
      </label>
    </div>
    <div class="checks">
      <label><input id="noSignals" type="checkbox"> 不输出 signals</label>
      <label><input id="noActions" type="checkbox"> 不输出 actions</label>
    </div>
    <button id="run" type="submit">开始回测</button>
  </form>
  <section><div id="status" class="status">等待运行</div><pre id="output">尚未运行。</pre></section>
</main>
<script>
const $ = id => document.getElementById(id);
function selectedStrategies() {
  return [...document.querySelectorAll("#strategyOptions input[type=checkbox]:checked")].map(input => input.value);
}
function updateStrategySummary() {
  const values=selectedStrategies(), all=document.getElementById("strategyAll"), options=document.querySelectorAll("#strategyOptions input[type=checkbox]");
  document.getElementById("strategy").value=values.join(" ");
  document.getElementById("strategyToggle").textContent=values.length ? (values.length === options.length ? `已选择全部策略（${values.length}个）` : `已选择 ${values.length} 个策略`) : "请选择策略";
  all.checked=values.length>0 && values.length===options.length;
  all.indeterminate=values.length>0 && !all.checked;
}
let activeStrategyCategory="全部";
let strategyCatalog=[];
function renderStrategyCategories() {
  const categories=["全部", ...new Set(strategyCatalog.map(item=>item.category))];
  const container=document.getElementById("strategyCategories"); container.replaceChildren();
  categories.forEach(category=>{
    const button=document.createElement("button"); button.type="button"; button.className="strategy-category";
    button.textContent=category; button.classList.toggle("active", category===activeStrategyCategory);
    button.addEventListener("click",event=>{
      event.preventDefault();
      event.stopPropagation();
      activeStrategyCategory=category;
      renderStrategyCategories();
      filterStrategyOptions();
      document.getElementById("strategyMenu").hidden=false;
    });
    container.append(button);
  });
}
function filterStrategyOptions() {
  document.querySelectorAll("#strategyOptions .strategy-option").forEach(label=>{
    const visible=activeStrategyCategory==="全部" || label.dataset.category===activeStrategyCategory;
    label.hidden=!visible;
    label.style.display=visible ? "flex" : "none";
  });
}
function renderStrategies(catalog) {
  strategyCatalog=catalog; renderStrategyCategories();
  const container=document.getElementById("strategyOptions"); container.replaceChildren();
  catalog.forEach(item=>{
    const name=item.name, category=item.category;
    const label=document.createElement("label"); label.className="strategy-option";
    label.dataset.category=category;
    const input=document.createElement("input"); input.type="checkbox"; input.value=name; input.checked=name==="ma5ma10";
    input.addEventListener("change",updateStrategySummary); label.append(input,document.createTextNode(name)); container.append(label);
  });
  filterStrategyOptions();
  updateStrategySummary();
}
async function loadStrategies() {
  const response=await fetch("/api/strategies"), data=await response.json();
  renderStrategies(data.strategies || []);
}
document.getElementById("strategyToggle").addEventListener("click",()=>{const menu=document.getElementById("strategyMenu");menu.hidden=!menu.hidden;});
document.getElementById("strategySelectAll").addEventListener("click",()=>{document.querySelectorAll("#strategyOptions input[type=checkbox]").forEach(input=>input.checked=true);updateStrategySummary();});
document.getElementById("strategyClearAll").addEventListener("click",()=>{document.querySelectorAll("#strategyOptions input[type=checkbox]").forEach(input=>input.checked=false);updateStrategySummary();});
document.getElementById("strategyAll").addEventListener("change",event=>{document.querySelectorAll("#strategyOptions input[type=checkbox]").forEach(input=>input.checked=event.target.checked);updateStrategySummary();});
document.addEventListener("click",event=>{
  const picker=document.getElementById("strategyPicker");
  if(!picker.contains(event.target)) document.getElementById("strategyMenu").hidden=true;
});
loadStrategies().catch(error=>{document.getElementById("strategyToggle").textContent="策略列表加载失败";console.error(error);});
$("form").addEventListener("submit", async (event) => {
  event.preventDefault();
  $("run").disabled = true; $("status").textContent = "正在回测，请等待..."; $("status").className = "status";
  $("output").textContent = "";
  const value = id => $(id).value.trim();
  const payload = {
    strategy: value("strategy"), timeframe: value("timeframe"), direction: value("direction"), volume: value("volume"), symbols: value("symbols"), adapter: value("adapter"), level: value("level"),
    start: value("start"), end: value("end"), data_source: value("dataSource"), data_dir: value("dataDir"),
    prefix: value("prefix"), out_dir: value("outDir"), warmup_bars: value("warmupBars"), config_dir: value("configDir"),
    no_signals: $("noSignals").checked, no_actions: $("noActions").checked
  };
  try {
    const response = await fetch("/api/run", {method: "POST", headers: {"Content-Type": "application/json"}, body: JSON.stringify(payload)});
    const result = await response.json(); $("output").textContent = result.output || result.error || "";
    $("status").textContent = result.ok ? "回测完成" : "回测失败";
    $("status").className = "status " + (result.ok ? "ok" : "bad");
  } catch (error) { $("status").textContent = "请求失败"; $("status").className = "status bad"; $("output").textContent = error; }
  finally { $("run").disabled = false; }
});
</script>
</body>
</html>"""


def _tokens(value: str) -> list[str]:
    return value.split()


def build_command(payload: dict) -> list[str]:
    command = [sys.executable, str(BASE / "run_backtest.py")]
    strategies = _tokens(str(payload.get("strategy", "")))
    if strategies:
        command += ["--strategy", *strategies]
    symbols = _tokens(str(payload.get("symbols", "")))
    if symbols:
        command += ["--symbols", *symbols]
    for key, option in (("adapter", "--adapter"), ("level", "--level"), ("timeframe", "--timeframe"),
                        ("volume", "--volume"),
                        ("direction", "--direction"), ("start", "--start"),
                        ("end", "--end"), ("data_dir", "--data-dir"), ("data_source", "--data-source"),
                        ("prefix", "--prefix"), ("out_dir", "--out-dir"), ("warmup_bars", "--warmup-bars"),
                        ("config_dir", "--config-dir")):
        value = str(payload.get(key, "")).strip()
        if value:
            command += [option, value]
    if payload.get("no_signals"):
        command.append("--no-signals")
    if payload.get("no_actions"):
        command.append("--no-actions")
    return command


def decode_output(raw: bytes) -> str:
    """Decode child-process output from UTF-8 or the Windows Chinese code page."""
    for encoding in ("utf-8", "gb18030", "cp936"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _json_default(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    try:
        if value != value:  # NaN
            return None
    except Exception:
        pass
    return str(value)


def _clean_json(value):
    """Convert pandas NaN/Infinity values to JSON null recursively."""
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, dict):
        return {key: _clean_json(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_clean_json(item) for item in value]
    return value


def _safe_output_file(value: str) -> Path:
    """Resolve a report path while preventing access outside output/."""
    candidate = (OUTPUT / value).resolve()
    root = OUTPUT.resolve()
    if candidate != root and root not in candidate.parents:
        raise ValueError("文件路径必须位于 strategy/output 目录内")
    if not candidate.is_file():
        raise FileNotFoundError(candidate)
    return candidate


def _prefix_file(folder: str, prefix: str, suffix: str) -> Path:
    return OUTPUT / folder / f"{suffix}_{prefix}.csv"


def discover_runs() -> list[dict]:
    import pandas as pd

    runs = []
    capital_dir = OUTPUT / "capital"
    for path in sorted(capital_dir.glob("capital_daily_*.csv"), key=lambda p: p.stat().st_mtime, reverse=True):
        prefix = path.stem.removeprefix("capital_daily_")
        timeframe = "1d"
        metadata = {}
        metadata_path = OUTPUT / "runs" / f"{prefix}.json"
        if metadata_path.is_file():
            try:
                metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
                timeframe = metadata.get("timeframe", "1d")
            except Exception:
                pass
        try:
            frame = pd.read_csv(path, usecols=["symbol", "strategy"])
            symbols = sorted(frame["symbol"].dropna().astype(str).unique().tolist())
            strategies = sorted(frame["strategy"].dropna().astype(str).unique().tolist())
        except Exception:
            try:
                frame = pd.read_csv(path, usecols=["symbol"])
                symbols = sorted(frame["symbol"].dropna().astype(str).unique().tolist())
            except Exception:
                symbols = []
            strategies = [str(item) for item in metadata.get("strategies", [])]
        runs.append({
            "prefix": prefix,
            "symbols": symbols,
            "strategies": strategies,
            "capital": str(path.relative_to(OUTPUT)).replace("\\", "/"),
            "performance": f"performance/performance_report_{prefix}.xlsx" if (OUTPUT / "performance" / f"performance_report_{prefix}.xlsx").is_file() else None,
            "evaluation": f"evaluation/evaluation_report_{prefix}.xlsx" if (OUTPUT / "evaluation" / f"evaluation_report_{prefix}.xlsx").is_file() else None,
            "overview": f"capital/symbol_core_overview_{prefix}.xlsx" if (OUTPUT / "capital" / f"symbol_core_overview_{prefix}.xlsx").is_file() else None,
            "timeframe": timeframe,
        })
    return runs


def _chart_indicators(bars, strategy: str | None) -> dict:
    """Build indicators for the selected strategy, aligned with chart bars."""
    import numpy as np
    import pandas as pd

    if bars is None or bars.empty or not strategy or strategy in {"__all__", "all"}:
        return {"main": [], "sub": []}
    name = str(strategy).lower()
    close = bars["close"].astype(float)
    high = bars["high"].astype(float)
    low = bars["low"].astype(float)
    volume = bars["volume"].astype(float)
    main: list[dict] = []
    sub: list[dict] = []

    def values(series):
        return [None if pd.isna(v) or not np.isfinite(float(v)) else float(v) for v in series]

    if any(token in name for token in ("ma5ma10", "double_ma", "ma_crossover", "multiperiod_ma", "guppy", "hull_ma", "kama")):
        main.extend([
            {"name": "MA5", "color": "#f59e0b", "values": values(close.rolling(5).mean())},
            {"name": "MA10", "color": "#8b5cf6", "values": values(close.rolling(10).mean())},
        ])
    elif any(token in name for token in ("boll", "keltner", "channel", "donchian", "supertrend")):
        mid = close.rolling(20).mean()
        std = close.rolling(20).std(ddof=0)
        main.extend([
            {"name": "中轨", "color": "#f59e0b", "values": values(mid)},
            {"name": "上轨", "color": "#2563eb", "values": values(mid + 2 * std)},
            {"name": "下轨", "color": "#2563eb", "values": values(mid - 2 * std)},
        ])
    elif "kdj" in name or "stochastic" in name:
        lowest = low.rolling(9).min()
        highest = high.rolling(9).max()
        rsv = ((close - lowest) / (highest - lowest).replace(0, np.nan) * 100).fillna(50)
        k = rsv.ewm(alpha=1 / 3, adjust=False).mean()
        d = k.ewm(alpha=1 / 3, adjust=False).mean()
        j = 3 * k - 2 * d
        sub.extend([
            {"name": "K", "color": "#2563eb", "values": values(k)},
            {"name": "D", "color": "#f59e0b", "values": values(d)},
            {"name": "J", "color": "#dc2626", "values": values(j)},
        ])
    elif "macd" in name:
        dif = close.ewm(span=12, adjust=False).mean() - close.ewm(span=26, adjust=False).mean()
        dea = dif.ewm(span=9, adjust=False).mean()
        sub.extend([
            {"name": "DIF", "color": "#2563eb", "values": values(dif)},
            {"name": "DEA", "color": "#f59e0b", "values": values(dea)},
            {"name": "MACD", "color": "#7c3aed", "values": values((dif - dea) * 2)},
        ])
    elif any(token in name for token in ("rsi", "zscore", "mean_reversion")):
        change = close.diff()
        up = change.clip(lower=0).rolling(14).mean()
        down = (-change.clip(upper=0)).rolling(14).mean()
        rsi = 100 - 100 / (1 + up / down.replace(0, np.nan))
        sub.append({"name": "RSI14", "color": "#2563eb", "values": values(rsi)})
    elif "cci" in name:
        typical = (high + low + close) / 3
        mean = typical.rolling(20).mean()
        deviation = typical.rolling(20).apply(lambda x: np.mean(np.abs(x - np.mean(x))), raw=True)
        cci = (typical - mean) / (0.015 * deviation.replace(0, np.nan))
        sub.append({"name": "CCI20", "color": "#7c3aed", "values": values(cci)})
    elif any(token in name for token in ("volume", "money_flow", "obv", "mfi", "turnover", "open_interest")):
        direction = np.sign(close.diff().fillna(0))
        obv = (direction * volume).cumsum()
        sub.append({"name": "OBV", "color": "#059669", "values": values(obv)})
    elif "atr" in name:
        previous = close.shift(1)
        true_range = pd.concat([high - low, (high - previous).abs(), (low - previous).abs()], axis=1).max(axis=1)
        sub.append({"name": "ATR14", "color": "#dc2626", "values": values(true_range.rolling(14).mean())})
    return {"main": main, "sub": sub}


def load_result(prefix: str, symbol: str, data_dir: str | None = None,
                timeframe: str | None = None, strategy: str | None = None) -> dict:
    import pandas as pd

    capital_path = OUTPUT / "capital" / f"capital_daily_{prefix}.csv"
    if not capital_path.is_file():
        raise FileNotFoundError(f"没有找到回测结果: {capital_path.name}")
    capital_all = pd.read_csv(capital_path)
    capital_all["symbol"] = capital_all["symbol"].astype(str)
    capital_symbol = capital_all[capital_all["symbol"] == symbol].copy()
    if capital_symbol.empty:
        raise ValueError(f"结果 {prefix} 中没有品种 {symbol}")
    strategy_filter = None if not strategy or strategy in {"__all__", "all"} else str(strategy)
    capital = capital_symbol
    if strategy_filter and "strategy" in capital.columns:
        capital = capital[capital["strategy"].astype(str) == strategy_filter].copy()
    if capital.empty:
        raise ValueError(f"结果 {prefix} / {symbol} 中没有策略 {strategy_filter}")

    def read_rows(folder: str, stem: str, limit: int | None = 1000, selected_strategy: str | None = strategy_filter):
        path = OUTPUT / folder / f"{stem}_{prefix}.csv"
        if not path.is_file():
            return []
        frame = pd.read_csv(path)
        if "symbol" in frame.columns:
            frame = frame[frame["symbol"].astype(str) == symbol]
        if selected_strategy and "strategy" in frame.columns:
            frame = frame[frame["strategy"].astype(str) == selected_strategy]
        if limit is not None:
            frame = frame.tail(limit)
        return frame.to_dict(orient="records")

    # Keep the complete symbol trade history so early buy/sell markers are not
    # silently dropped when a run contains more than 500 trades.
    all_trades = read_rows("trades", "trades", None, None)
    trades = read_rows("trades", "trades", None)
    signals = read_rows("signals", "signals", 1000)
    daily = read_rows("daily/symbol", "daily_pnl", 2000)

    data_path = Path(data_dir or "../data/future_price")
    if not data_path.is_absolute():
        data_path = (BASE / data_path).resolve()
    metadata_path = OUTPUT / "runs" / f"{prefix}.json"
    if not timeframe and metadata_path.is_file():
        try:
            timeframe = json.loads(metadata_path.read_text(encoding="utf-8")).get("timeframe")
        except Exception:
            timeframe = None
    timeframe = str(timeframe or "1d").lower()
    if timeframe not in {"15m", "1h", "1d"}:
        timeframe = "1d"
    candles = []
    bars = None
    indicator_data = {"main": [], "sub": []}
    h5_path = data_path / f"{symbol}.h5"
    if h5_path.is_file():
        from data_sources.hdf5_index import HDF5IndexSource
        bars = HDF5IndexSource(str(data_path)).load(symbol)
        if bars is not None and not bars.empty:
            bars = bars.copy()
            bars["tradingday"] = pd.to_datetime(bars["tradingday"])
            if timeframe == "1d":
                bars["period"] = bars["tradingday"]
            else:
                bars["period"] = bars.index.floor("15min" if timeframe == "15m" else "1h")
            daily_bars = bars.groupby("period", sort=True).agg(
                bar_time=("period", "first"), tradingday=("tradingday", "first"), open=("open", "first"),
                high=("high", "max"), low=("low", "min"),
                close=("close", "last"), volume=("volume", "sum")
            ).reset_index(drop=True)
            first_day = pd.to_datetime(capital["date"]).min()
            last_day = pd.to_datetime(capital["date"]).max()
            daily_bars = daily_bars[(daily_bars["tradingday"] >= first_day) & (daily_bars["tradingday"] <= last_day)]
            # 权益曲线使用完整回测区间；K 线也返回完整区间，避免固定
            # tail(1200) 让日线图从 2021 年附近开始而不是从回测起点开始。
            candles = daily_bars.to_dict(orient="records")
            indicator_data = _chart_indicators(daily_bars, strategy_filter)

    equity_frame = capital.groupby("date", as_index=False).agg(
        initial_capital=("initial_capital", "sum"),
        realized_pnl=("realized_pnl", "sum"),
        unrealized_pnl=("unrealized_pnl", "sum"),
        equity=("equity", "sum"),
        daily_pnl=("daily_pnl", "sum"),
        net_value=("net_value", "sum"),
    ).sort_values("date")
    equity = equity_frame.tail(3000).to_dict(orient="records")
    equity_values = equity_frame["equity"].astype(float)
    running_max = equity_values.cummax()
    drawdown = ((equity_values - running_max) / running_max).min() * 100
    final_equity = float(equity_values.iloc[-1])
    profits = [float(row.get("profit", 0) or 0) for row in trades]
    metrics = {
        "initial_capital": float(equity_frame["initial_capital"].iloc[0]),
        "final_equity": final_equity,
        "total_pnl": final_equity - float(equity_frame["initial_capital"].iloc[0]),
        "max_drawdown_pct": float(drawdown),
        "trade_count": len(trades),
        "win_rate_pct": (sum(p > 0 for p in profits) / len(profits) * 100) if profits else 0,
    }
    from performance.funding_metrics import account_metrics, margin_metrics
    run_period = {}
    completion_path = OUTPUT / "runs" / f"{prefix}.complete.json"
    if completion_path.is_file():
        run_period = json.loads(completion_path.read_text(encoding="utf-8"))
    metrics.update(account_metrics(equity_frame, run_period.get("start"), run_period.get("end")))
    from core.data_loader import load_future_list
    future_list = load_future_list()
    info = future_list.loc[symbol] if symbol in future_list.index else {}
    selected_strategies = set(capital["strategy"].astype(str))
    metrics.update(margin_metrics(read_rows("actions", "actions", None, None), bars,
                   metrics["initial_capital"], float(info.get("point", float("nan"))),
                   float(info.get("margin_rate", float("nan"))), selected_strategies,
                   run_period.get("start") or equity_frame.date.iloc[0],
                   run_period.get("end") or str(equity_frame.date.iloc[-1])+" 23:59:59"))
    if metrics.get("margin_available") and metrics["max_margin_amount"] > 0:
        metrics["drawdown_margin_multiple"] = metrics["max_drawdown_amount"]/metrics["max_margin_amount"]
        metrics["annual_profit_margin_multiple"] = metrics["simple_annual_pnl"]/metrics["max_margin_amount"]
    strategy_summary = []
    if "strategy" in capital_symbol.columns:
        for strategy_name, group in capital_symbol.groupby("strategy", sort=True):
            group_equity = group.groupby("date", as_index=False).agg(
                initial_capital=("initial_capital", "sum"), equity=("equity", "sum")
            ).sort_values("date")
            values = group_equity["equity"].astype(float)
            running = values.cummax()
            group_drawdown = ((values - running) / running).min() * 100 if len(values) else 0
            group_trades = [row for row in all_trades if str(row.get("strategy", "")) == str(strategy_name)]
            group_profits = [float(row.get("profit", 0) or 0) for row in group_trades]
            group_initial = float(group_equity["initial_capital"].iloc[0]) if len(group_equity) else 0
            group_final = float(values.iloc[-1]) if len(values) else group_initial
            strategy_summary.append({
                "strategy": str(strategy_name),
                "trade_count": len(group_trades),
                "total_pnl": group_final - group_initial,
                "max_drawdown_pct": float(group_drawdown),
                "win_rate_pct": (sum(p > 0 for p in group_profits) / len(group_profits) * 100) if group_profits else 0,
            })
    return {"prefix": prefix, "symbol": symbol, "strategy": strategy_filter or "__all__", "timeframe": timeframe,
            "candles": candles, "indicators": indicator_data, "equity": equity, "signals": signals, "trades": trades, "daily": daily,
            "metrics": metrics, "strategy_summary": strategy_summary}


def load_report(path_value: str, sheet: str | None = None) -> dict:
    from openpyxl import load_workbook

    path = _safe_output_file(path_value)
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = workbook.sheetnames
    selected = sheet if sheet in sheets else (sheets[0] if sheets else None)
    rows = []
    if selected:
        for row in workbook[selected].iter_rows(max_row=201, values_only=True):
            rows.append([_json_default(cell) if cell is not None else "" for cell in row])
    workbook.close()
    return {"file": path_value, "sheets": sheets, "sheet": selected, "rows": rows}


VIEWER_HTML = r"""
<section id="viewer" style="margin-top:20px">
  <h2>回测结果查看器</h2>
  <p class="hint">读取 output 中已有的 CSV/XLSX，不改变原有结果文件。</p>
  <div class="grid">
    <label>回测批次<select id="runSelect"></select></label>
    <label style="display:none">品种<select id="symbolSelect"></select></label>
    <label>行情目录<input id="viewerDataDir" value="../data/future_price"></label>
  </div>
  <div class="viewer-layout">
    <div class="viewer-main">
      <button id="loadResult" type="button">读取结果</button>
      <div id="metrics" class="grid" style="margin-top:18px"></div>
      <section><h3>策略汇总</h3><div id="strategySummary"></div></section>
      <div class="chart-grid" style="margin-top:18px">
        <section><h3>K线与买卖点</h3><div class="chart-tools"><button type="button" onclick="zoomChart('kline',0.7)">放大</button><button type="button" onclick="zoomChart('kline',1.4)">缩小</button><button type="button" onclick="resetChart('kline')">全部</button></div><canvas id="kline" width="1100" height="420" style="width:100%;height:auto;border:1px solid #e5e7eb"></canvas><div id="klineTip" class="chart-tip"></div></section>
        <section><h3>权益曲线</h3><div class="chart-tools"><button type="button" onclick="zoomChart('equity',0.7)">放大</button><button type="button" onclick="zoomChart('equity',1.4)">缩小</button><button type="button" onclick="resetChart('equity')">全部</button></div><canvas id="equity" width="1100" height="420" style="width:100%;height:auto;border:1px solid #e5e7eb"></canvas><div id="equityTip" class="chart-tip"></div></section>
      </div>
      <section><h3>交易记录</h3><div id="tradesTable"></div></section>
      <section><h3>详细报告</h3><div id="reports"></div><div id="reportTable" style="overflow:auto;margin-top:12px"></div></section>
    </div>
    <aside class="viewer-side">
      <h3 style="margin-top:0">策略</h3>
      <select id="resultStrategySelect" style="display:none"></select>
      <div id="resultStrategyList" class="symbol-list" tabindex="0" aria-label="策略列表"></div>
      <h3 style="margin-top:0">品种</h3>
      <p class="hint">点击品种，或先点击列表后使用键盘 ↑ ↓ ← → 切换。</p>
      <div id="symbolList" class="symbol-list" tabindex="0" aria-label="品种列表"></div>
    </aside>
  </div>
</section>
<script>
const runs = [];
const esc = value => String(value ?? "").replace(/[&<>"']/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;","\"":"&quot;","'":"&#39;"}[c]));
async function refreshRuns() {
  const response = await fetch("/api/runs"); const data = await response.json();
  runs.splice(0, runs.length, ...(data.runs || []));
  const select = document.getElementById("runSelect");
  select.innerHTML = runs.map((r,i) => `<option value="${esc(r.prefix)}">${esc(r.prefix)} [${esc(r.timeframe || "1d")}] (${r.symbols.join(", ")})</option>`).join("");
  updateSymbols();
}
function currentRun() { return runs.find(r => r.prefix === document.getElementById("runSelect").value); }
function selectSymbol(symbol, load=true) {
  const select=document.getElementById("symbolSelect");
  if (![...select.options].some(option => option.value === symbol)) return;
  select.value=symbol;
  document.querySelectorAll(".symbol-item").forEach(item => item.classList.toggle("active", item.dataset.symbol === symbol));
  if (load) loadResult();
}
function selectResultStrategy(strategy, load=true) {
  const select=document.getElementById("resultStrategySelect");
  if (![...select.options].some(option => option.value === strategy)) return;
  select.value=strategy;
  document.querySelectorAll("#resultStrategyList .symbol-item").forEach(item => item.classList.toggle("active", item.dataset.strategy === strategy));
  if (load) loadResult();
}
function updateSymbols() {
  const r=currentRun(), symbols=r?.symbols||[], strategies=r?.strategies||[], select=document.getElementById("symbolSelect"), list=document.getElementById("symbolList"), strategySelect=document.getElementById("resultStrategySelect"), strategyList=document.getElementById("resultStrategyList");
  select.innerHTML=symbols.map(s=>`<option value="${esc(s)}">${esc(s)}</option>`).join("");
  list.innerHTML=symbols.map(s=>`<div class="symbol-item" tabindex="-1" data-symbol="${esc(s)}">${esc(s)}</div>`).join("");
  list.querySelectorAll(".symbol-item").forEach(item => item.addEventListener("click",()=>{selectSymbol(item.dataset.symbol);list.focus();}));
  const resultStrategies=["__all__", ...strategies];
  strategySelect.innerHTML=resultStrategies.map(s=>`<option value="${esc(s)}">${s==="__all__"?"全部策略":esc(s)}</option>`).join("");
  strategyList.innerHTML=resultStrategies.map(s=>`<div class="symbol-item" tabindex="-1" data-strategy="${esc(s)}">${s==="__all__"?"全部策略":esc(s)}</div>`).join("");
  strategyList.querySelectorAll(".symbol-item").forEach(item => item.addEventListener("click",()=>{selectResultStrategy(item.dataset.strategy);strategyList.focus();}));
  if (resultStrategies.length) selectResultStrategy("__all__", false);
  if (symbols.length) selectSymbol(symbols[0], false);
}
function moveSymbol(delta) {
  const select=document.getElementById("symbolSelect"), count=select.options.length;
  if (!count) return;
  const next=(select.selectedIndex+delta+count)%count;
  selectSymbol(select.options[next].value);
}
document.getElementById("runSelect").addEventListener("change", updateSymbols);
document.getElementById("symbolList").addEventListener("keydown", event => {
  if (["ArrowUp","ArrowLeft"].includes(event.key)) { event.preventDefault(); moveSymbol(-1); }
  if (["ArrowDown","ArrowRight"].includes(event.key)) { event.preventDefault(); moveSymbol(1); }
});
 const chartState={kline:{rows:[],signals:[],trades:[],mainIndicators:[],subIndicators:[],timeframe:"1d",start:0,end:0},equity:{rows:[],start:0,end:0},viewStart:0,viewEnd:1,crosshair:null};
function syncChartWindows(){[chartState.kline,chartState.equity].forEach(state=>{const n=state.rows.length;state.start=Math.max(0,Math.floor(chartState.viewStart*n));state.end=Math.min(n,Math.max(state.start+1,Math.ceil(chartState.viewEnd*n)));});}
function range(state){return state.rows.slice(state.start,state.end||state.rows.length);}
function globalToLocalX(fraction,state){return (fraction-chartState.viewStart)/(chartState.viewEnd-chartState.viewStart||1);}
function drawBoth(){drawKline();drawEquity();}
function nice(v){return Number(v).toLocaleString(undefined,{maximumFractionDigits:2});}
function movingAverage(rows,period){return rows.map((_,i)=>i+1<period?null:rows.slice(i-period+1,i+1).reduce((sum,r)=>sum+Number(r.close),0)/period);}
function drawKline(){
  const state=chartState.kline, rows=range(state), c=document.getElementById("kline"),ctx=c.getContext("2d"),w=c.width,h=c.height;ctx.clearRect(0,0,w,h);
  if(!rows.length){ctx.fillText("未找到对应 H5 行情数据",20,30);return;}
  const main=state.mainIndicators||[],sub=state.subIndicators||[],hasSub=sub.length>0;
  const left=62,right=72,top=24,priceBottom=Math.round(h*(hasSub?.55:.68)),subTop=hasSub?Math.round(h*.61):0,subBottom=hasSub?Math.round(h*.80):0,volumeTop=hasSub?Math.round(h*.86):priceBottom,volumeBottom=h-45,plotW=w-left-right;
  const lo=Math.min(...rows.map(r=>Number(r.low))),hi=Math.max(...rows.map(r=>Number(r.high))),span=hi-lo||1,maxVol=Math.max(...rows.map(r=>Number(r.volume)||0),1);
  const px=i=>left+(i+.5)*plotW/rows.length, py=v=>priceBottom-((v-lo)/span)*(priceBottom-top), vy=v=>volumeBottom-(v/maxVol)*(volumeBottom-volumeTop),cw=Math.max(2,Math.min(16,plotW/rows.length*.65));
  ctx.font="12px sans-serif";ctx.strokeStyle="#e5e7eb";ctx.fillStyle="#6b7280";ctx.strokeRect(left,top,plotW,volumeBottom-top);
  for(let n=0;n<=5;n++){const y=top+n*(priceBottom-top)/5;ctx.beginPath();ctx.moveTo(left,y);ctx.lineTo(w-right,y);ctx.stroke();ctx.fillText(nice(hi-n*span/5),w-right+8,y+4);}
  rows.forEach((r,i)=>{const x=px(i),o=Number(r.open),cl=Number(r.close),hv=Number(r.high),lv=Number(r.low),up=cl>=o,color=up?"#ef4444":"#10b981";ctx.strokeStyle=color;ctx.fillStyle=color;ctx.beginPath();ctx.moveTo(x,py(lv));ctx.lineTo(x,py(hv));ctx.stroke();ctx.fillRect(x-cw/2,Math.min(py(o),py(cl)),cw,Math.max(1,Math.abs(py(o)-py(cl))));ctx.fillRect(x-cw/2,vy(Number(r.volume)||0),cw,volumeBottom-vy(Number(r.volume)||0));if(i%Math.max(1,Math.floor(rows.length/8))===0)ctx.fillText(String(r.bar_time||r.tradingday).replace("T"," ").slice(0,16),x-36,volumeBottom+18);});
  const drawSeries=(items,scaleY,labelY)=>items.forEach((item,index)=>{const vals=item.values||[];ctx.strokeStyle=item.color||["#2563eb","#f59e0b","#dc2626"][index%3];ctx.lineWidth=1.8;ctx.beginPath();let started=false;rows.forEach((_,i)=>{const v=vals[state.start+i];if(v==null)return;const x=px(i),y=scaleY(v);started?ctx.lineTo(x,y):ctx.moveTo(x,y);started=true;});ctx.stroke();ctx.fillStyle=ctx.strokeStyle;ctx.fillText(item.name,left+8+index*48,labelY+16);});
  drawSeries(main,py,top);
  if(hasSub){const subValues=sub.flatMap(item=>item.values.slice(state.start,state.end||state.rows.length).filter(v=>v!=null));let subLo=Math.min(...subValues),subHi=Math.max(...subValues);if(!Number.isFinite(subLo)||!Number.isFinite(subHi)){subLo=0;subHi=100;}if(subLo===subHi){subLo-=1;subHi+=1;}const subSpan=subHi-subLo,sy=v=>subBottom-((v-subLo)/subSpan)*(subBottom-subTop);ctx.strokeStyle="#e5e7eb";ctx.strokeRect(left,subTop,plotW,subBottom-subTop);for(let n=0;n<=2;n++){const y=subTop+n*(subBottom-subTop)/2;ctx.beginPath();ctx.moveTo(left,y);ctx.lineTo(w-right,y);ctx.stroke();ctx.fillStyle="#6b7280";ctx.fillText(nice(subHi-n*subSpan/2),w-right+8,y+4);}drawSeries(sub,sy,subTop);ctx.fillStyle="#6b7280";ctx.fillText("指标",left,subTop-6);}
  const keyForTime=value=>{const text=String(value||"").replace("T"," ");if(state.timeframe==="1d")return text.slice(0,10);if(state.timeframe==="15m")return text.slice(0,16);return text.slice(0,13);};
  const byTime={};rows.forEach((r,i)=>{const key=keyForTime(r.bar_time||r.tradingday);if(byTime[key]==null)byTime[key]=i;});
  // Use actual execution timestamps/prices from trades, not signal timestamps.
  // Signals are generated at the close and market_0905 orders fill the next day.
  (state.trades||[]).forEach(t=>{
    const direction=String(t.direction||"").toLowerCase(), openDate=keyForTime(t.open_time), closeDate=keyForTime(t.close_time);
    [[openDate,Number(t.open_price),"开"+(direction==="long"?"多":"空"),direction==="long"],[closeDate,Number(t.close_price),"平"+(direction==="long"?"多":"空"),direction!=="long"]].forEach(item=>{
      const i=byTime[item[0]];if(i==null||!Number.isFinite(item[1]))return;const x=px(i),y=py(item[1]),up=Boolean(item[3]);ctx.fillStyle=up?"#16a34a":"#dc2626";ctx.beginPath();ctx.moveTo(x,y+(up?-16:16));ctx.lineTo(x-7,y+(up?0:0));ctx.lineTo(x+7,y+(up?0:0));ctx.closePath();ctx.fill();ctx.fillText(item[2],x-13,y+(up?-21:28));
    });
  });
  if(chartState.crosshair!=null){const local=globalToLocalX(chartState.crosshair,state);if(local>=0&&local<=1){const x=left+local*plotW;ctx.strokeStyle="#64748b";ctx.setLineDash([5,4]);ctx.beginPath();ctx.moveTo(x,top);ctx.lineTo(x,volumeBottom);ctx.stroke();ctx.setLineDash([]);}}
  ctx.fillStyle="#6b7280";ctx.fillText("VOL",left,volumeBottom+36);ctx.fillText("价格",w-right+8,top-7);ctx.fillText("主图指标：MA5 / MA10",left+120,top+16);
  const legend=main.length?main.map(item=>item.name).join(" / "):"无";
  ctx.fillStyle="#ffffff";ctx.fillRect(left+105,top,260,24);ctx.fillStyle="#6b7280";ctx.fillText("主图指标："+legend,left+120,top+16);
 }
 function drawEquity(){
  const state=chartState.equity,rows=range(state),c=document.getElementById("equity"),ctx=c.getContext("2d"),w=c.width,h=c.height;ctx.clearRect(0,0,w,h);
  if(!rows.length){ctx.fillText("暂无权益数据",20,30);return;}
  const left=70,right=80,top=28,bottom=h-45,plotW=w-left-right,values=rows.map(r=>Number(r.equity)),lo=Math.min(...values),hi=Math.max(...values),span=hi-lo||1,px=i=>left+(i+.5)*plotW/rows.length,py=v=>bottom-((v-lo)/span)*(bottom-top);
  ctx.strokeStyle="#e5e7eb";ctx.fillStyle="#6b7280";ctx.strokeRect(left,top,plotW,bottom-top);for(let n=0;n<=5;n++){const y=top+n*(bottom-top)/5;ctx.beginPath();ctx.moveTo(left,y);ctx.lineTo(w-right,y);ctx.stroke();ctx.fillText(nice(hi-n*span/5),5,y+4);}ctx.strokeStyle="#2563eb";ctx.lineWidth=2;ctx.beginPath();rows.forEach((r,i)=>{const x=px(i),y=py(Number(r.equity));i?ctx.lineTo(x,y):ctx.moveTo(x,y);});ctx.stroke();if(chartState.crosshair!=null){const local=globalToLocalX(chartState.crosshair,state);if(local>=0&&local<=1){const x=left+local*plotW;ctx.strokeStyle="#64748b";ctx.setLineDash([5,4]);ctx.beginPath();ctx.moveTo(x,top);ctx.lineTo(x,bottom);ctx.stroke();ctx.setLineDash([]);}}ctx.fillStyle="#2563eb";rows.forEach((r,i)=>{if(i%Math.max(1,Math.floor(rows.length/8))===0)ctx.fillText(String(r.date).slice(0,10),px(i)-25,bottom+18);});ctx.fillStyle="#6b7280";ctx.fillText("权益",w-right+8,top-7);
}
function zoomChart(name,factor){const span=chartState.viewEnd-chartState.viewStart;if(!span)return;const target=Math.max(20/Math.max(chartState[name].rows.length,1),Math.min(1,span*factor)),center=chartState.crosshair==null?(chartState.viewStart+chartState.viewEnd)/2:chartState.crosshair;chartState.viewStart=Math.max(0,center-target/2);chartState.viewEnd=Math.min(1,chartState.viewStart+target);chartState.viewStart=Math.max(0,chartState.viewEnd-target);syncChartWindows();drawBoth();}
function resetChart(name){chartState.viewStart=0;chartState.viewEnd=1;chartState.crosshair=null;syncChartWindows();drawBoth();}
 function attachChartHover(name){const c=document.getElementById(name),tip=document.getElementById(name+"Tip");c.addEventListener("wheel",e=>{e.preventDefault();zoomChart(name,e.deltaY<0?.75:1.35);},{passive:false});c.addEventListener("mousemove",e=>{const s=chartState[name],rect=c.getBoundingClientRect(),x=(e.clientX-rect.left)*c.width/rect.width,left=name==="kline"?62:70,right=name==="kline"?72:80,local=Math.max(0,Math.min(1,(x-left)/(c.width-left-right)));chartState.crosshair=chartState.viewStart+local*(chartState.viewEnd-chartState.viewStart);const rows=s.rows,idx=Math.max(0,Math.min(rows.length-1,Math.round(chartState.crosshair*(rows.length-1)))),r=rows[idx];if(!r)return;tip.textContent=name==="kline"?`时间 ${String(r.bar_time||r.tradingday).replace("T"," ").slice(0,16)} | O ${nice(r.open)} H ${nice(r.high)} L ${nice(r.low)} C ${nice(r.close)} V ${nice(r.volume)}`:`日期 ${r.date} | 权益 ${nice(r.equity)} | 当日盈亏 ${nice(r.daily_pnl)}`;tip.style.display="block";drawBoth();});c.addEventListener("mouseleave",()=>{tip.style.display="none";chartState.crosshair=null;drawBoth();});}
function table(rows, columns) { if(!rows?.length)return "<p>暂无数据</p>"; return `<table style="border-collapse:collapse;width:100%;font-size:12px"><thead><tr>${columns.map(c=>`<th style="border:1px solid #ddd;padding:5px;text-align:left">${esc(c)}</th>`).join("")}</tr></thead><tbody>${rows.slice(-100).reverse().map(r=>`<tr>${columns.map(c=>`<td style="border:1px solid #ddd;padding:5px">${esc(r[c])}</td>`).join("")}</tr>`).join("")}</tbody></table>`; }
async function loadResult() {
  const prefix=document.getElementById("runSelect").value,symbol=document.getElementById("symbolSelect").value,strategy=document.getElementById("resultStrategySelect").value,dataDir=document.getElementById("viewerDataDir").value,run=currentRun(),timeframe=run?.timeframe||"1d";
  if(!prefix||!symbol)return; const response=await fetch(`/api/result?prefix=${encodeURIComponent(prefix)}&symbol=${encodeURIComponent(symbol)}&strategy=${encodeURIComponent(strategy)}&data_dir=${encodeURIComponent(dataDir)}&timeframe=${encodeURIComponent(timeframe)}`); const d=await response.json(); if(d.error){alert(d.error);return;}
  const m=d.metrics, money=v=>Number(v).toLocaleString('zh-CN',{minimumFractionDigits:2,maximumFractionDigits:2})+' 元', pct=v=>Number(v).toFixed(4)+'%', times=v=>Number(v).toFixed(6)+' 倍';
  const metricRows=[['初始本金',money(m.initial_capital)],['最终权益',money(m.final_equity)],['总收益（含未平仓浮盈亏）',money(m.total_pnl)],['总收益率／初始本金',pct(m.total_return_pct)],['最终权益／初始本金',times(m.final_equity_multiple)],['最大回撤金额（日末盯市）',money(m.max_drawdown_amount)],['最大回撤／初始本金',pct(m.max_drawdown_initial_pct)],['标准最大回撤率／历史峰值',pct(-m.max_drawdown_pct)],['单利年化收益额（365日折算）',money(m.simple_annual_pnl)],['单利年化收益率／初始本金',pct(m.simple_annual_return_pct)],['单利年化利润／初始本金',times(m.simple_annual_profit_multiple)],['最大估算保证金占用',m.margin_available?money(m.max_margin_amount):'无法估算'],['最大占用／初始本金',m.margin_available?pct(m.max_margin_initial_pct):'—'],['最大回撤／最大估算保证金',m.drawdown_margin_multiple!=null?times(m.drawdown_margin_multiple):'—'],['单利年化利润／最大估算保证金',m.annual_profit_margin_multiple!=null?times(m.annual_profit_margin_multiple):'—'],['交易次数',m.trade_count],['胜率',pct(m.win_rate_pct)]];
  metricRows.push(['相对初始本金最大亏损',money(m.max_initial_loss_amount)],['相对初始本金最大亏损率',pct(m.max_initial_loss_pct)],['最大名义持仓额（非保证金）',m.margin_available?money(m.max_notional_amount):'—']);
  document.getElementById("metrics").innerHTML=metricRows.map(x=>`<div style="background:#eff6ff;padding:12px;border-radius:8px"><small>${x[0]}</small><br><strong>${x[1]}</strong></div>`).join('')+`<p class="hint">回撤金额取历史权益峰值至后续谷值，包含初始本金基准；按初始本金归一化与标准峰值回撤率分开显示。相对初始本金最大亏损则取初始本金至最低权益的差额，与峰谷回撤不同。年化按完整区间日历天数单利折算。保证金仅为固定比例估算，保证金倍数不等于账户亏损／收益倍数，也不表示用该金额即可运行。${esc(m.margin_note||'')}</p>`;
  document.getElementById("strategySummary").innerHTML=table(d.strategy_summary,['strategy','trade_count','total_pnl','max_drawdown_pct','win_rate_pct']);
  chartState.kline.rows=d.candles||[];chartState.kline.signals=d.signals||[];chartState.kline.trades=d.trades||[];chartState.kline.timeframe=d.timeframe||timeframe;chartState.kline.mainIndicators=d.indicators?.main||[];chartState.kline.subIndicators=d.indicators?.sub||[];chartState.equity.rows=d.equity||[];chartState.viewStart=0;chartState.viewEnd=1;chartState.crosshair=null;syncChartWindows();drawBoth();document.getElementById("tradesTable").innerHTML=table(d.trades,['strategy','trade_date','direction','open_time','close_time','open_price','close_price','volume','profit','close_reason']);
  const r=currentRun();document.getElementById("reports").innerHTML=[['绩效报告',r.performance],['深度评价',r.evaluation],['品种概览',r.overview]].filter(x=>x[1]).map(x=>`<button type="button" class="reportBtn" data-file="${esc(x[1])}">${esc(x[0])}</button>`).join(" ");
  document.querySelectorAll('.reportBtn').forEach(b=>b.addEventListener('click',()=>loadReport(b.dataset.file)));
}
function matrixTable(rows){if(!rows?.length)return "<p>暂无数据</p>";const heads=rows[0].map((v,i)=>v===null||v===undefined||v===""?`列${i+1}`:v);return `<table style="border-collapse:collapse;width:100%;font-size:12px"><thead><tr>${heads.map(h=>`<th style="border:1px solid #ddd;padding:6px;text-align:left;background:#f1f5f9;position:sticky;top:0">${esc(h)}</th>`).join("")}</tr></thead><tbody>${rows.slice(1).map(r=>`<tr>${heads.map((_,i)=>`<td style="border:1px solid #ddd;padding:6px">${esc(r[i])}</td>`).join("")}</tr>`).join("")}</tbody></table>`;}
async function loadReport(file,sheet) { const url=`/api/report?file=${encodeURIComponent(file)}${sheet?`&sheet=${encodeURIComponent(sheet)}`:""}`;const d=await (await fetch(url)).json();if(d.error){alert(d.error);return;}document.getElementById("reportTable").innerHTML=`<p>文件：${esc(d.file)}　工作表：<select id="sheetSelect">${d.sheets.map(s=>`<option ${s===d.sheet?'selected':''}>${esc(s)}</option>`).join("")}</select></p>`+matrixTable(d.rows);document.getElementById("sheetSelect").addEventListener('change',e=>loadReport(file,e.target.value));}
attachChartHover("kline");attachChartHover("equity");document.getElementById("loadResult").addEventListener("click",loadResult);refreshRuns();
</script>
"""


class Handler(BaseHTTPRequestHandler):
    def _send(self, status: int, body: str, content_type: str = "text/html; charset=utf-8") -> None:
        encoded = body.encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(encoded)))
        self.end_headers()
        self.wfile.write(encoded)

    def do_GET(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        if parsed.path == "/":
            self._send(200, HTML.replace("</body>", VIEWER_HTML + "</body>"))
        elif parsed.path == "/api/strategies":
            self._send(200, json.dumps({"strategies": discover_strategy_catalog()}, ensure_ascii=False), "application/json; charset=utf-8")
        elif parsed.path == "/api/runs":
            self._send(200, json.dumps({"runs": discover_runs()}, ensure_ascii=False, default=_json_default), "application/json; charset=utf-8")
        elif parsed.path == "/api/result":
            from urllib.parse import parse_qs
            query = parse_qs(parsed.query)
            try:
                result = load_result(
                    query.get("prefix", [""])[0],
                    query.get("symbol", [""])[0],
                    query.get("data_dir", [None])[0],
                    query.get("timeframe", [None])[0],
                    query.get("strategy", [None])[0],
                )
                self._send(200, json.dumps(_clean_json(result), ensure_ascii=False, default=_json_default, allow_nan=False), "application/json; charset=utf-8")
            except Exception as exc:
                self._send(400, json.dumps({"error": str(exc)}, ensure_ascii=False), "application/json; charset=utf-8")
        elif parsed.path == "/api/report":
            from urllib.parse import parse_qs
            query = parse_qs(parsed.query)
            try:
                report = load_report(query.get("file", [""])[0], query.get("sheet", [None])[0])
                self._send(200, json.dumps(_clean_json(report), ensure_ascii=False, default=_json_default, allow_nan=False), "application/json; charset=utf-8")
            except Exception as exc:
                self._send(400, json.dumps({"error": str(exc)}, ensure_ascii=False), "application/json; charset=utf-8")
        else:
            self._send(404, "Not found", "text/plain; charset=utf-8")

    def do_POST(self) -> None:  # noqa: N802
        if urlparse(self.path).path != "/api/run":
            self._send(404, json.dumps({"error": "Not found"}), "application/json")
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            command = build_command(payload)
            result = subprocess.run(command, cwd=BASE, capture_output=True)
            stdout = decode_output(result.stdout or b"")
            stderr = decode_output(result.stderr or b"")
            output = stdout + ("\n" + stderr if stderr else "")
            body = json.dumps({"ok": result.returncode == 0, "output": output, "returncode": result.returncode}, ensure_ascii=False)
            self._send(200, body, "application/json; charset=utf-8")
        except Exception as exc:  # pragma: no cover - UI error response
            self._send(400, json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False), "application/json; charset=utf-8")

    def log_message(self, format: str, *args) -> None:
        return


if __name__ == "__main__":
    print(f"CTA web console (local): http://127.0.0.1:{PORT}")
    print(f"Bind address: {HOST}:{PORT}  (LAN clients use this computer's IPv4 address)")
    print(f"Working directory: {BASE}")
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
